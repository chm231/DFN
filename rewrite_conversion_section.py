from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


REPORT = Path(__file__).resolve().parent / "docs" / "DFM샘플데이터_파이프라인_데모보고서_정리본.docx"


def set_font(run, size=10.5, color="202020", bold=False):
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Malgun Gothic")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Malgun Gothic")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def replace_paragraph(paragraph, text, style=None):
    paragraph._p.clear_content()
    if style:
        paragraph.style = style
    run = paragraph.add_run(text)
    set_font(run)


def remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)


def main():
    doc = Document(REPORT)
    section_start = next(
        i for i, p in enumerate(doc.paragraphs)
        if p.text.startswith("2. 전처리·변환")
    )
    section_end = next(
        i for i, p in enumerate(doc.paragraphs[section_start + 1:], section_start + 1)
        if p.text == "전역 절리군"
    )
    targets = doc.paragraphs[section_start + 1:section_end]

    replacement = [
        ("Normal", "DFM Export를 파이프라인 표준 입력 trace_dataset_3d.h5/.csv로 변환하는 단계다. 변환은 절리선의 3차원 좌표와 전역 절리군, 관측창 정보 및 검열 상태를 일관된 스키마로 기록한다."),
        ("Normal", "python scripts/convert_dfm_export_to_trace_dataset.py \\\n  --dfm-dir demodata/DFM_Export/DFM_Export --outdir demo_output/dfm_demo"),
        ("Normal", "주요 옵션은 --grid(점유격자 셀 크기, 기본 0.05 m), --edge-tol(censoring 판정 거리, 기본 0.15 m), --cluster-cut(절리군 군집 절단각, 기본 30°)이다."),
        ("Normal", "변환 규약:"),
        ("List Bullet", "터널축 정의 — 터널축 x-hat은 12개 면 중심 궤적을 선형 적합한 방향으로 정한다(world x-hat = [+0.9995, -0.0216, +0.0214]). 실측 막장면 법선은 굴진방향에 대해 평균 8.2° 기울어 있으므로, 법선 평균을 축으로 사용하지 않아 단면의 x 방향 드리프트를 방지한다."),
        ("List Bullet", "전역 절리군 재분류 — 면별 DS 번호는 전역 set으로 직접 사용하지 않고, 평균 법선의 축간각 30° 기준 군집화로 모든 관측을 전역 set으로 재분류한다."),
        ("List Bullet", "관측면적과 대표 관측창 — 관측면적은 면별 점유격자 합으로 795.4 m²를 사용한다. 단일 대표 관측창은 면적 중앙값 면의 convex hull로 정한다."),
        ("List Bullet", "끝점 검열 — 절리선 끝점이 해당 면 hull 경계에서 0.15 m 이내이면 관측창 경계에 의한 censoring으로 판정한다."),
        ("List Bullet", "미지 항목 — 반경 참값이 없으므로 radius_m은 NaN으로 기록하고, fracture_id는 파일 내 연번으로 부여한다."),
        ("Normal", "출력은 trace_dataset_3d.h5/.csv, 조건부 생성 및 P32 단계 입력용 dfn_export_for_python.h5(관측 법선·대표 관측창), 그리고 변환 결과 검증용 진단 파일 2종으로 구성된다."),
    ]

    for paragraph, (style, text) in zip(targets, replacement):
        replace_paragraph(paragraph, text, style)
    for paragraph in reversed(targets[len(replacement):]):
        remove_paragraph(paragraph)

    doc.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    main()

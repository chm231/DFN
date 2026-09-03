from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
REPORT = ROOT / "docs" / "DFM샘플데이터_파이프라인_데모보고서_정리본.docx"
ISO = ROOT / "storage" / "output" / "pipeline_demo_laxemar" / "export" / "domain_dfn_3d_iso.png"
SIDE = ROOT / "storage" / "output" / "pipeline_demo_laxemar" / "export" / "domain_dfn_3d_side.png"


def set_font(run, size, color="202020", bold=False):
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Malgun Gothic")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Malgun Gothic")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table):
    widths = (4680, 4680)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_panel(cell, image_path, label, description):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    picture = p.add_run().add_picture(str(image_path), width=Inches(3.02))
    picture._inline.docPr.set("descr", description)
    cap = cell.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(0)
    r = cap.add_run(label)
    set_font(r, 8.5, color="6B7280")


def main():
    if not REPORT.exists() or not ISO.exists() or not SIDE.exists():
        raise FileNotFoundError("보고서 또는 3D 도메인 그림 파일을 찾을 수 없습니다.")
    doc = Document(REPORT)
    doc.add_page_break()
    heading = doc.add_paragraph(style="Heading 1")
    heading.paragraph_format.space_after = Pt(8)
    heading.add_run("부록 A. 3차원 해석 도메인 시각화")
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(8)
    r = intro.add_run("관측 절리, 미관측 절리, 굴착 터널 및 터널벽 절리선의 공간적 관계를 등각뷰와 측면뷰로 제시한다.")
    set_font(r, 10.5)
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table)
    add_panel(table.cell(0, 0), ISO, "(a) 등각뷰", "해석 도메인 DFN 등각뷰")
    add_panel(table.cell(0, 1), SIDE, "(b) 측면뷰 (x-z 평면)", "해석 도메인 DFN 측면뷰")
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(0)
    r = caption.add_run("그림 A-1. 터널 굴착 구간과 DFN 해석 도메인의 3차원 구성")
    set_font(r, 9, color="6B7280")
    doc.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    main()

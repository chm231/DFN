from __future__ import annotations

import base64
import re
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "docs" / "DFM샘플데이터_파이프라인_데모보고서.html"
OUT = ROOT / "docs" / "DFM샘플데이터_파이프라인_데모보고서_정리본.docx"
ASSET_DIR = ROOT / "docs" / "_dfm_report_assets"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "6B7280"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"


class Node:
    def __init__(self, tag: str, attrs=None):
        self.tag = tag
        self.attrs = dict(attrs or [])
        self.children = []

    def text(self) -> str:
        return "".join(c if isinstance(c, str) else c.text() for c in self.children).strip()


class SourceParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node("root")
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        node = Node(tag.lower(), attrs)
        self.stack[-1].children.append(node)
        if tag.lower() not in {"img", "br", "meta", "link", "hr"}:
            self.stack.append(node)

    def handle_startendtag(self, tag, attrs):
        self.stack[-1].children.append(Node(tag.lower(), attrs))

    def handle_endtag(self, tag):
        tag = tag.lower()
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                self.stack = self.stack[:i]
                break

    def handle_data(self, data):
        self.stack[-1].children.append(data)


def set_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Malgun Gothic")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Malgun Gothic")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = tc_mar.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}")
            tc_mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Malgun Gothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption.font.name = "Malgun Gothic"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("DFM 실측 샘플 데이터 파이프라인 데모 보고서")
    set_font(r, size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("2026.08.20  |  ")
    set_font(r, size=8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)


def add_cover(doc):
    doc.add_paragraph().paragraph_format.space_after = Pt(72)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("TECHNICAL DEMONSTRATION REPORT")
    set_font(r, size=10, color=BLUE, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("DFM 실측 샘플 데이터\n파이프라인 데모 보고서")
    set_font(r, size=26, color=INK, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(42)
    r = p.add_run("실측 막장면 자료 기반 DFN 복원 · 조건화 · 안정성 해석 입력 생성")
    set_font(r, size=12, color=MUTED)
    table = doc.add_table(rows=3, cols=2)
    set_table_geometry(table, [2200, 7160])
    for i, (label, value) in enumerate([
        ("작성일", "2026-08-20"),
        ("입력 데이터", "실측 터널 막장면 12개 및 DFM Viewer 절리 매핑 추출본"),
        ("파이프라인", "handoffv1 - hybrid k_r, 해석식 P32, 응집 연결 복원, remove-and-resample 조건화"),
    ]):
        left, right = table.rows[i].cells
        set_cell_shading(left, LIGHT_BLUE)
        for cell, text, bold in ((left, label, True), (right, value, False)):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            rr = p.add_run(text)
            set_font(rr, size=9.5, color=INK, bold=bold)
    doc.add_page_break()


def direct_children(node, tag):
    return [c for c in node.children if isinstance(c, Node) and c.tag == tag]


def descendants(node, tag):
    found = []
    for child in node.children:
        if isinstance(child, Node):
            if child.tag == tag:
                found.append(child)
            found.extend(descendants(child, tag))
    return found


def add_table_from_html(doc, node):
    rows = []
    for tr in descendants(node, "tr"):
        cells = [c for c in tr.children if isinstance(c, Node) and c.tag in {"th", "td"}]
        if cells:
            rows.append(cells)
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    base, remainder = divmod(9360, cols)
    set_table_geometry(table, [base + (1 if i < remainder else 0) for i in range(cols)])
    for i, html_row in enumerate(rows):
        for j, html_cell in enumerate(html_row):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(re.sub(r"\s+", " ", html_cell.text()))
            set_font(r, size=8.7, color=INK, bold=(i == 0))
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_image(doc, node, index):
    src = node.attrs.get("src", "")
    alt = node.attrs.get("alt", f"결과 그림 {index}")
    if not src.startswith("data:image/"):
        return
    match = re.match(r"data:image/([^;]+);base64,(.*)", src, flags=re.S)
    if not match:
        return
    ASSET_DIR.mkdir(exist_ok=True)
    image_path = ASSET_DIR / f"figure_{index}.{match.group(1).replace('jpeg', 'jpg')}"
    image_path.write_bytes(base64.b64decode(match.group(2)))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(image_path), width=Inches(6.05))
    cap = doc.add_paragraph(style="Figure Caption")
    cap.add_run(f"그림 {index}. {alt}")


def add_text_paragraph(doc, text, style=None):
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return
    p = doc.add_paragraph(style=style)
    p.paragraph_format.widow_control = True
    r = p.add_run(text)
    set_font(r, size=10.5, color="202020")


def walk_content(doc, node, figures):
    for child in node.children:
        if isinstance(child, str):
            continue
        if child.tag in {"style", "script", "head", "title"}:
            continue
        if child.tag == "h1":
            continue
        if child.tag == "h2":
            p = doc.add_paragraph(style="Heading 1")
            p.add_run(child.text())
        elif child.tag == "h3":
            p = doc.add_paragraph(style="Heading 2")
            p.add_run(child.text())
        elif child.tag in {"p", "div", "section", "article", "body"}:
            imgs = direct_children(child, "img")
            tables = direct_children(child, "table")
            direct_text = "".join(c for c in child.children if isinstance(c, str)).strip()
            if direct_text:
                add_text_paragraph(doc, direct_text)
            if child.tag == "p" and not imgs and not tables:
                text = child.text()
                if text:
                    add_text_paragraph(doc, text)
            for img in imgs:
                figures[0] += 1
                add_image(doc, img, figures[0])
            for table in tables:
                add_table_from_html(doc, table)
            for nested in child.children:
                if isinstance(nested, Node) and nested.tag not in {"img", "table"}:
                    walk_content(doc, nested, figures)
        elif child.tag == "table":
            add_table_from_html(doc, child)
        elif child.tag == "img":
            figures[0] += 1
            add_image(doc, child, figures[0])
        elif child.tag in {"ul", "ol"}:
            for li in direct_children(child, "li"):
                p = doc.add_paragraph(style="List Bullet")
                r = p.add_run(li.text())
                set_font(r, size=10.5, color="202020")
        elif child.tag == "pre":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(child.text())
            r.font.name = "Consolas"
            r.font.size = Pt(8.5)
        else:
            walk_content(doc, child, figures)


def main():
    parser = SourceParser()
    parser.feed(SOURCE.read_text(encoding="utf-8"))
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    body = next((n for n in parser.root.children if isinstance(n, Node) and n.tag == "html"), parser.root)
    walk_content(doc, body, [0])
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
